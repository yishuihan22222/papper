"""Final cross-file verification of V7 input, results and reviewed Word output."""
from pathlib import Path
import hashlib
import json
import re
import importlib
import numpy as np
import pandas as pd
import pymupdf
from docx import Document

audit=importlib.import_module('136_audit_processing_units_v7')
OUT=audit.OUT
BASE=Path(r'E:\sci23')
FOLDER=BASE/'论文节点整改_V7_20261021'
QA=FOLDER/'排版与核验'


def main():
    design=audit.read_json(OUT/'00_design.json')
    assert audit.sha(OUT/'00_design.json')==audit.read_json(OUT/'design_lock.json')['design_sha256']
    assert all(audit.sha(p)==h for p,h in design['source_hashes'].items())
    assert all(audit.sha(OUT/p)==h for p,h in design['audit_hashes'].items())
    done=audit.read_json(OUT/'completion.json')
    assert done['completed'] and done['runs']==6912 and done['passed']==done['checks']==914
    rows=pd.read_csv(OUT/'01_all_246_node_decisions.csv')
    assert len(rows)==rows.node_id.nunique()==246
    assert rows.main_included.sum()==122 and rows.decision.eq('exclude_role').sum()==34
    assert rows.decision.str.startswith('quarantine_').sum()==87
    assert rows.decision.str.startswith('reference_').sum()==3
    assert not rows.independent_unit_verified.any()
    near=pd.read_csv(OUT/'05_all_nearby_or_homonym_pairs.csv')
    assert len(near)==179 and not near.confirmed_same_processing_unit.any()
    nn=pd.read_csv(OUT/'model_inputs/nodes.csv',float_precision='round_trip')
    ee=pd.read_csv(OUT/'model_inputs/edges.csv',float_precision='round_trip')
    assert len(nn)==122 and len(ee)==1624 and nn.node_id.tolist()==design['keep_ids']
    assert set(ee.Source_id)|set(ee.Target_id)<=set(nn.node_id)
    role=rows.set_index('node_id')
    assert nn.family.tolist()==role.loc[nn.node_id,'reviewed_family'].tolist()
    demand=audit.read_json(OUT/'model_inputs/retained_source_demands.json')
    resources=pd.read_csv(OUT/'14_resources.csv')
    for name,value in demand.items():
        assert value['node_ids']==nn.node_id.tolist()
        assert np.array(value['parts']).shape[1]==len(nn)
        observed=resources[resources.demand.eq(name)].task_denominator.to_numpy()
        assert np.max(abs(observed-np.array(value['parts']).sum()))<1e-7
    summary=pd.read_csv(OUT/'12_policy_summary.csv')
    main=summary[summary.scope.eq('candidate_corrected_family')]
    assert main.C_B_better.sum()==265 and main.C_B_equal.sum()==311 and main.C_B_worse.sum()==0
    assert np.max(abs(main.C_B_median_pp))<=1e-9
    word=BASE/'中文论文初稿_节点处理单元整改V7_20261021.docx'
    doc=Document(word)
    assert len(doc.tables)==16 and len(doc.inline_shapes)==11 and len(doc.sections)==3
    assert len(doc.element.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'))==21
    text='\n'.join(p.text for p in doc.paragraphs)
    assert all(x not in text for x in ['@figure','@table','@references','TODO','TBD'])
    assert all(x in text for x in ['122个处理场地候选','1624条','6912条轨迹','265次','311次','1.72—2.98'])
    assert re.search(r'\[20\]',text)
    previous=pymupdf.open(QA/'V7_核验_最终.pdf');pdf=pymupdf.open(QA/'V7_核验_交付.pdf')
    assert len(pdf)==len(previous)==43
    assert all(len(p.get_text().strip())>20 and '\ufffd' not in p.get_text() for p in pdf)
    changed=[i for i,(a,b) in enumerate(zip(previous,pdf),1) if a.get_pixmap().samples!=b.get_pixmap().samples]
    assert changed==[27,28]
    assert all(re.sub(r'\s','',a.get_text())==re.sub(r'\s','',b.get_text()) for a,b in zip(previous,pdf))
    v5=BASE/'中文论文初稿_全图美化与正文优化V5_20261019.docx'
    assert audit.sha(v5)=='a4eb07d3e7cfc7b2c5be8d4c011d35515ad8896f6c2c896f662a57a65ebb1b63'
    for p in (FOLDER/'图件').glob('*.png'):
        assert audit.sha(p)==audit.sha(BASE/'论文收口优化_V6_20261021/图件'/p.name)
    sources=audit.read_json(OUT/'17_source_registry.json')
    verified=audit.read_json(OUT/'18_source_body_checks.json')
    assert len(verified)==7 and all(x['body_verified'] for x in verified)
    assert all(audit.sha(r['local_path'])==r['sha256'] for r in sources['new_sources'] if r['downloaded'])
    record=dict(date='2026-10-21',status='passed',nodes_audited=246,retained_candidates=122,
        excluded_role=34,quarantined=87,hierarchy_references=3,physical_duplicate_merges=0,
        relation_pairs_reviewed=179,edges=1624,model_runs=6912,policy_conditions=2304,
        numerical_integrity_checks=914,cold_restarts=8,protected_files_unchanged=len(design['source_hashes']),
        original_V5_and_V6_unchanged=True,pages=43,all_pages_visually_reviewed=True,
        visual_review_pages=list(range(1,44)),last_layout_change_pages=[27,28],other_41_pages_pixel_identical=True,
        figures=11,tables=16,equations=21,references=20,main_tables=6,main_figures=4,
        archived_source_bodies_verified=7,one_source_unarchived='VAILOG_CURRENT TLS certificate mismatch; web tool reading retained as limited support only',
        all_retained_independently_verified=False,actual_real_capacity_replacements=0,
        renderer='Microsoft Word PDF export and PyMuPDF images; bundled LibreOffice unavailable as diagnosed in V6',
        word=str(word),word_sha256=audit.sha(word),verification_pdf=str(QA/'V7_核验_交付.pdf'))
    audit.dump(record,QA/'最终核验.json')
    paths=[word,BASE/'运行修订模型_V7.ps1',*OUT.rglob('*'),*FOLDER.rglob('*')]
    code=Path(__file__).parent
    paths += [code/f for f in ['136_audit_processing_units_v7.py','137_run_processing_unit_rectification_v7.py','138_package_processing_unit_v7.py','139_prepare_manuscript_v7.py','140_verify_v7_delivery.py']]
    manifest=[]
    for path in sorted(set(paths)):
        if not path.is_file() or path.name=='delivery_manifest.json':continue
        if '排版与核验' in str(path) and (path.suffix.lower()=='.png' or path.name.startswith('V7_核验') and path.name!='V7_核验_交付.pdf'):continue
        manifest.append(dict(path=str(path),bytes=path.stat().st_size,sha256=audit.sha(path)))
    audit.dump(dict(schema='v7-final-delivery',validation=record,files=manifest),OUT/'delivery_manifest.json')
    print(json.dumps(record,ensure_ascii=False,indent=2),flush=True)


if __name__=='__main__':main()
